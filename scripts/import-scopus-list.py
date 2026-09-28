#!/usr/bin/env python3
"""
import-scopus-list.py — build src/lib/scopus-sources.json from Elsevier's
Scopus source title list (the "Download the Source title list" xlsx on
Scopus's content coverage page).

Keeps only what the journal pages need to link each journal to its own
Scopus page and state its status in that list, keyed by every ISSN/eISSN
(8 characters, no hyphen):

    [source_id | null, code, detail | null]
      code "a"  active                      detail: first coverage year
      code "i"  inactive (coverage ended)   detail: last coverage year
      code "x"  discontinued by Scopus      detail: year of discontinuation
      code "p"  accepted, being added       detail: date of acceptance

A title absent from the list is not evidence that Scopus lacks it: the list
holds ~49,000 serial titles of the 50,000+ in Scopus, so the site never
says "not in Scopus".

Usage (openpyxl required: pip install openpyxl):
    python3 scripts/import-scopus-list.py <ext_list_<Mon>_<YYYY>.xlsx> [--as-of 2026-08]
"""
import json, re, sys
from pathlib import Path
import openpyxl

src = sys.argv[1]
as_of = sys.argv[sys.argv.index('--as-of') + 1] if '--as-of' in sys.argv else None
wb = openpyxl.load_workbook(src, read_only=True)

def sheet(prefix):
    return next(ws for ws in wb.worksheets if ws.title.startswith(prefix))

sources = sheet('Scopus Sources')
if not as_of:
    m = re.search(r'(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\w*\.?\s+(\d{4})', sources.title)
    months = 'Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec'.split()
    as_of = f"{m.group(2)}-{months.index(m.group(1)) + 1:02d}" if m else None

def issn(v):
    v = re.sub(r'[^0-9Xx]', '', str(v or '')).upper()
    return v.zfill(8) if 7 <= len(v) <= 8 else None

def years(coverage):
    ys = [int(y) for y in re.findall(r'\b(1[89]\d\d|20\d\d)\b', str(coverage or ''))]
    return (min(ys), max(ys)) if ys else (None, None)

# Year each source was discontinued, from the discontinued-titles sheet.
disc_year = {}
rows = sheet('Discontinued Titles').iter_rows(values_only=True)
next(rows); head = next(rows)
for r in rows:
    if r[0] is None or r[5] != 'Discontinuation':
        continue
    y = r[6]
    if isinstance(y, (int, float)):
        disc_year[int(r[0])] = max(disc_year.get(int(r[0]), 0), int(y))

# Preference when one ISSN maps to several sources (title changes).
RANK = {'a': 0, 'p': 1, 'i': 2, 'x': 3}
out = {}
def put(key, entry):
    if key and (key not in out or RANK[entry[1]] < RANK[out[key][1]]):
        out[key] = entry

rows = sources.iter_rows(values_only=True)
head = next(rows)
ix = {h: i for i, h in enumerate(head)}
n = 0
for r in rows:
    if r[0] is None:
        continue
    n += 1
    sid = int(r[ix['Sourcerecord ID']])
    first, last = years(r[ix['Coverage']])
    if str(r[ix['Titles Discontinued by Scopus']] or '').strip():
        entry = [sid, 'x', disc_year.get(sid, last)]
    elif r[ix['Active or Inactive']] == 'Active':
        entry = [sid, 'a', first]
    else:
        entry = [sid, 'i', last]
    for k in (issn(r[ix['ISSN']]), issn(r[ix['EISSN']])):
        put(k, entry)

accepted = 0
rows = sheet('Accepted Titles').iter_rows(values_only=True)
for r in rows:
    if r[0] in (None, 'Source Title') or (isinstance(r[0], str) and r[0].startswith('Notice')):
        continue
    d = r[3]
    date = d.strftime('%Y-%m-%d') if hasattr(d, 'strftime') else (str(d) if d else None)
    for k in (issn(r[1]), issn(r[2])):
        if k:
            accepted += 1
        put(k, [None, 'p', date])

path = Path(__file__).resolve().parent.parent / 'src/lib/scopus-sources.json'
path.write_text(json.dumps({'as_of': as_of, 'list': sources.title, 'd': dict(sorted(out.items()))}, separators=(',', ':')))
counts = {c: sum(1 for e in out.values() if e[1] == c) for c in 'aixp'}
print(f"{path.name}: {len(out)} ISSNs from {n} sources (as of {as_of}); by status {counts}; discontinued years {len(disc_year)}")
